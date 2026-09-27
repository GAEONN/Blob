"""Local document-to-handwriting renderer driven by a user's private glyph bank."""

import csv
from contextlib import ExitStack
import math
import random
import re
import unicodedata
from pathlib import Path

from PIL import Image, ImageDraw, ImageFilter, ImageFont

import handwriting_profile


TEXT_EXT = {".txt", ".md", ".docx"}
TABLE_EXT = {".csv", ".xlsx", ".xlsm"}
SUPPORTED_EXTS = TEXT_EXT | TABLE_EXT
PAGE_SIZE = (2550, 3300)
LINE_STEP = 88
INK_COLORS = {"negra": (28, 30, 42), "azul": (22, 40, 120)}
REPLACEMENTS = {
    "“": '"', "”": '"', "„": '"', "‘": "'", "’": "'", "´": "'",
    "–": "-", "—": "-", "…": "...", "×": "x", "•": "-", "·": ".",
    "\u00a0": " ", "\t": "    ", "\u200b": "",
}


class HandwritingError(ValueError):
    pass


def normalize(text):
    text = unicodedata.normalize("NFC", str(text))
    for old, new in REPLACEMENTS.items():
        text = text.replace(old, new)
    return text


def read_document(path, mode=None):
    path = Path(path)
    extension = path.suffix.lower()
    if extension not in SUPPORTED_EXTS:
        raise HandwritingError(f"Unsupported file type: {extension or 'no extension'}")
    if extension in (".txt", ".md"):
        content = path.read_text(encoding="utf-8-sig", errors="replace")
        detected = "texto"
    elif extension == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise HandwritingError("DOCX support is missing. Run Blob's installer to add it.") from exc
        document = Document(str(path))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            parts.extend("  |  ".join(cell.text.strip() for cell in row.cells)
                         for row in table.rows)
        content = "\n".join(parts)
        detected = "texto"
    elif extension == ".csv":
        with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as stream:
            sample = stream.read(4096)
            stream.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            rows = list(csv.reader(stream, dialect))
        content = "\n".join("  |  ".join(cell.strip() for cell in row) for row in rows)
        detected = "tabla"
    else:
        try:
            import openpyxl
        except ImportError as exc:
            raise HandwritingError("Excel support is missing. Run Blob's installer to add it.") from exc
        workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
        rows = []
        for sheet in workbook.worksheets:
            rows.append([sheet.title])
            rows.extend(["" if value is None else str(value) for value in row]
                        for row in sheet.iter_rows(values_only=True))
        workbook.close()
        content = "\n".join("  |  ".join(cell.strip() for cell in row) for row in rows)
        detected = "tabla"
    if mode == "tabla" and detected == "texto":
        content = "\n".join("  |  ".join(line.split("\t")) for line in content.splitlines())
    return normalize(content), mode or detected


def missing_glyphs(text, glyphs):
    return sorted({char for char in normalize(text) if not char.isspace() and char not in glyphs})


def _measure(text, glyphs, x_height):
    total = 0.0
    for char in text:
        if char.isspace():
            total += x_height * .76
        else:
            options = glyphs.get(char)
            if options:
                total += sum(item["w"] for item in options) / len(options) + x_height * .14
    return total


def _wrap(text, glyphs, x_height, max_width):
    out, line = [], ""
    for word in text.split():
        if _measure(word, glyphs, x_height) > max_width:
            if line:
                out.append(line)
                line = ""
            chunk = ""
            for char in word:
                if chunk and _measure(chunk + char, glyphs, x_height) > max_width:
                    out.append(chunk)
                    chunk = char
                else:
                    chunk += char
            line = chunk
            continue
        candidate = (line + " " + word).strip()
        if line and _measure(candidate, glyphs, x_height) > max_width:
            out.append(line)
            line = word
        else:
            line = candidate
    if line:
        out.append(line)
    return out


def _layout_lines(text, glyphs, x_height, max_width):
    layout = []
    for original in normalize(text).splitlines():
        line = original.strip()
        if not line:
            layout.append(("blank", ""))
        elif line.startswith("#"):
            layout.extend(("heading", chunk) for chunk in _wrap(line.lstrip("# "), glyphs, x_height, max_width))
        elif line.isupper() and len(line) <= 60:
            layout.extend(("heading", chunk) for chunk in _wrap(line, glyphs, x_height, max_width))
        else:
            layout.extend(("body", chunk) for chunk in _wrap(line, glyphs, x_height, max_width))
    while layout and layout[-1][0] == "blank":
        layout.pop()
    return layout


def _new_page(page_number, ink_color):
    image = Image.new("RGB", PAGE_SIZE, (250, 249, 246))
    draw = ImageDraw.Draw(image)
    width, height = PAGE_SIZE
    for y in range(300, height - 100, LINE_STEP):
        draw.line((0, y, width, y), fill=(209, 222, 239), width=2)
    draw.line((300, 0, 300, height), fill=(232, 177, 177), width=3)
    draw.text((width - 390, 130), f"{page_number}", fill=(155, 164, 174),
              font=ImageFont.truetype("arial.ttf", 25) if Path(r"C:\Windows\Fonts\arial.ttf").exists()
              else ImageFont.load_default())
    return image


def _write_line(page, text, glyphs, rng, baseline, ink_color, heading=False):
    x_height = handwriting_profile.TARGET_X_HEIGHT * (1.08 if heading else 1.0)
    left, right = 350, PAGE_SIZE[0] - 150
    width = _measure(text, glyphs, x_height)
    x = max(left, (PAGE_SIZE[0] - width) / 2) if heading else left + rng.uniform(-4, 7)
    start_x = x
    slant = rng.uniform(-.004, .004)
    for char in text:
        if char.isspace():
            x += x_height * rng.uniform(.62, .95)
            continue
        options = glyphs.get(char)
        if not options:
            x += x_height * .45
            continue
        glyph = rng.choice(options)
        base_scale = rng.uniform(.965, 1.035) * (1.05 if heading else 1.0)
        w = max(1, round(glyph["w"] * base_scale))
        h = max(1, round(glyph["h"] * base_scale))
        mask = glyph["image"].resize((w, h), Image.Resampling.LANCZOS)
        angle = rng.gauss(0, 1.0 if heading else 1.45)
        mask = mask.rotate(angle, resample=Image.Resampling.BICUBIC, expand=True)
        base = glyph["base"] * base_scale + (mask.height - h) / 2
        y = baseline - base + slant * (x - start_x) + rng.gauss(0, .8)
        # Paste only the small glyph mask. A transparent full-page layer for
        # every line caused needless allocation and made longer jobs sluggish.
        page.paste(ink_color, (round(x), round(y)), mask)
        x += w + x_height * rng.uniform(.08, .22)
        if x >= right:
            break


def render_document(path, profile_home=handwriting_profile.PROFILE_HOME, mode=None, ink="negra",
                    title=None, output_root=None, preview=False, progress=None, seed=None):
    profile = handwriting_profile.load_profile(profile_home)
    if not profile:
        raise HandwritingError("Create a local handwriting profile before writing a task.")
    glyphs = profile["glyphs"]
    text, detected_mode = read_document(path, mode)
    missing = missing_glyphs(text, glyphs)
    output_title = normalize(title if title is not None else Path(path).stem.upper())
    missing += [char for char in missing_glyphs(output_title, glyphs) if char not in missing]
    if missing:
        shown = " ".join(repr(char) for char in missing[:12])
        more = " …" if len(missing) > 12 else ""
        raise HandwritingError("Your profile is missing these characters: " + shown + more)
    if not text.strip():
        raise HandwritingError("The selected document is empty.")
    if ink not in INK_COLORS:
        raise HandwritingError("Choose black or blue ink.")
    x_height = handwriting_profile.TARGET_X_HEIGHT
    max_width = PAGE_SIZE[0] - 500
    lines = _layout_lines(text, glyphs, x_height, max_width)
    if output_title:
        lines.insert(0, ("heading", output_title))
        lines.insert(1, ("blank", ""))
    if not lines:
        raise HandwritingError("There is no text to write.")

    root = Path(output_root or (Path.home() / "Documents" / "Tareas a mano"))
    safe_name = re.sub(r"[^\w .-]+", "_", Path(path).stem, flags=re.UNICODE).strip(" .") or "Handwritten task"
    folder = root / (safe_name + " handwritten")
    folder.mkdir(parents=True, exist_ok=True)
    page_capacity = max(1, (PAGE_SIZE[1] - 440) // LINE_STEP)
    limit = min(2 * page_capacity, len(lines)) if preview else len(lines)
    page_count = max(1, math.ceil(limit / page_capacity))
    rng = random.Random(seed)
    color = INK_COLORS[ink]
    pages = []
    for page_index in range(page_count):
        page = _new_page(page_index + 1, color)
        start = page_index * page_capacity
        stop = min(limit, start + page_capacity)
        for row, (kind, line) in enumerate(lines[start:stop]):
            if kind == "blank":
                continue
            baseline = 355 + (row + 1) * LINE_STEP + rng.uniform(-1.3, 1.3)
            _write_line(page, line, glyphs, rng, baseline, color, heading=(kind == "heading"))
        page = page.filter(ImageFilter.GaussianBlur(.18))
        page_path = folder / f"page_{page_index + 1:02d}.jpg"
        page.save(page_path, quality=94, optimize=True)
        pages.append(page_path)
        if progress:
            progress(page_index + 1, page_count, f"Writing page {page_index + 1} of {page_count}")

    if preview:
        return {"folder": folder, "previews": [str(path) for path in pages], "pdf": None,
                "mode": detected_mode, "pages": page_count}
    pdf_path = folder / (safe_name + ".pdf")
    with ExitStack() as stack:
        first = stack.enter_context(Image.open(pages[0]))
        remainder = [stack.enter_context(Image.open(path)) for path in pages[1:]]
        first.save(pdf_path, format="PDF", save_all=True, append_images=remainder, resolution=300)
    return {"folder": folder, "previews": [], "pdf": str(pdf_path),
            "mode": detected_mode, "pages": page_count}
